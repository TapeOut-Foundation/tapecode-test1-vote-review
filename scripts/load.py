import openpyxl,collections
import sys
P=sys.argv[1] if len(sys.argv)>1 else 'TapeCode_Test1_NFT设计大赛_投票公示明细.xlsx'
wb=openpyxl.load_workbook(P,data_only=True)
def detail(n):
    return [r for r in wb[n].iter_rows(min_row=6,values_only=True) if r[0]]
def summ(n):
    return [r for r in wb[n].iter_rows(min_row=6,values_only=True) if r[0] is not None and r[1]]
